"""File -> text. No OCR: PDFs must have a text layer (scanned PDFs give no text)."""

import io
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}


class UnsupportedFileType(ValueError):
    pass


def extract_sections(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    """Return (page, text) sections. Page numbers for PDFs, None for other formats."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return [(None, _docx(data))]
    if ext in {".txt", ".md", ".csv"}:
        return [(None, data.decode("utf-8-sig", errors="replace"))]
    raise UnsupportedFileType(f"Unsupported file type {ext or '(none)'}; use {', '.join(sorted(SUPPORTED_EXTENSIONS))}")


def _pdf(data: bytes) -> list[tuple[int | None, str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return [(number, page.extract_text() or "") for number, page in enumerate(reader.pages, start=1)]


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)
